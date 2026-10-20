import re
from openpyxl import load_workbook


POSITION_RE = re.compile(r"^\d+\.\d+\.\d+$")

# Rexel broj: običan broj od minimum 5 cifara
REXEL_ARTICLE_RE = re.compile(r"\b\d{5,}\b")

# Würth broj: W + cifre, dozvoljavamo razmake, ali uzimamo samo W+cifre
WURTH_ARTICLE_RE = re.compile(r"\bW\s*(\d{5,})\b", re.IGNORECASE)


def position_to_tuple(position):
    """
    '1.20.017' -> (1, 20, 17)
    '1.20.0'   -> (1, 20, 0)
    """
    parts = str(position).strip().split(".")
    return tuple(int(part) for part in parts)


def is_position_in_range(position, start_position=None, end_position=None):
    current = position_to_tuple(position)

    if start_position:
        start = position_to_tuple(start_position)
        if current < start:
            return False

    if end_position:
        end = position_to_tuple(end_position)
        if current > end:
            return False

    return True


def extract_article_number(value):
    """
    Iz ćelije izvlači samo article number.

    Primeri:
    'W090632660   100' -> 'W090632660'
    'W0975259125'      -> 'W0975259125'
    '2035402'          -> '2035402'
    '2035402  1'       -> '2035402'
    'X'                -> None
    """

    text = str(value).strip()

    if not text:
        return None

    if text.upper() == "X":
        return None

    # Prvo proveri Würth broj
    wurth_match = WURTH_ARTICLE_RE.search(text)
    if wurth_match:
        digits = wurth_match.group(1)
        return f"W{digits}"

    # Zatim Rexel broj
    rexel_match = REXEL_ARTICLE_RE.search(text)
    if rexel_match:
        return rexel_match.group(0)

    return None


def parse_excel(file_path, column="A", sheet_name=None, start_position=None, end_position=None):
    """
    Čita Excel red po red.

    Pozicija:
        1.20.017

    Artikli ispod:
        2035402
        W0975259125
        W090632660   100

    X i prazno se preskaču.
    """

    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    current_position = None
    current_position_allowed = False
    tasks = []

    for row in range(1, ws.max_row + 1):
        raw_value = ws[f"{column}{row}"].value

        if raw_value is None:
            continue

        value = str(raw_value).strip()

        if not value:
            continue

        if value.upper() == "X":
            continue

        # Pozicija, npr. 1.20.017
        if POSITION_RE.match(value):
            current_position = value

            current_position_allowed = is_position_in_range(
                position=current_position,
                start_position=start_position,
                end_position=end_position
            )

            continue

        article_number = extract_article_number(value)

        if article_number:
            if current_position is None:
                print(f"[WARN] Red {row}: broj {article_number} nema poziciju iznad.")
                continue

            if not current_position_allowed:
                continue

            tasks.append({
                "position": current_position,
                "article_number": article_number,
                "raw_value": value,
                "row": row
            })

    return tasks